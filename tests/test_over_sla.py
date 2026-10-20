import contextlib
import datetime as dt
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, call, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import generate_report as report
import run_interactive as interactive
from pptx import Presentation


class OverSlaTests(unittest.TestCase):
    def setUp(self):
        report.set_report_tz(report.DEFAULT_REPORT_TZ)
        report.set_holiday_dates(None)
        self.addCleanup(report.set_holiday_dates, None)
        self.repo = report.parse_repo_url("https://dev.azure.com/org/project/_git/repo")
        self.team = [("Test Author", "author@example.test")]
        self.since = dt.date(2026, 5, 1)
        self.until = dt.date(2026, 7, 31)
        self.client = Mock(spec=report.AdoClient)
        self.client.completed_prs_for_user.return_value = [
            {
                "pullRequestId": pr_id,
                "closedDate": "2026-05-15T18:00:00Z",
                "targetRefName": target,
            }
            for pr_id, target in [
                (1, "refs/heads/main"),
                (2, "refs/heads/main"),
                (3, "refs/heads/main"),
                (4, "refs/heads/feature/test"),
                (5, "refs/heads/main"),
                (6, "refs/heads/main"),
            ]
        ]
        self.client.pr_last_commit_date.return_value = dt.date(2026, 5, 4)
        approvals = {
            1: dt.date(2026, 5, 4),  # 0
            2: dt.date(2026, 5, 5),  # 1
            3: dt.date(2026, 5, 6),  # 2
            4: dt.date(2026, 5, 7),  # 3
            5: dt.date(2026, 5, 8),  # 4
            6: dt.date(2026, 5, 1),  # approval before last commit
        }
        self.client.pr_last_approval_date.side_effect = (
            lambda repo, pr_id: approvals[pr_id]
        )
        self.client.pr_work_item_ids.return_value = [123]

    def extract(self, over_sla=False):
        return report.extract_rows(
            self.client, [self.repo], self.team, self.since, self.until,
            over_sla=over_sla,
        )[self.repo]

    def test_default_keeps_all_rows(self):
        rows = self.extract()
        self.assertEqual([row.business_days for row in rows], [0, 1, 2, 3, 4, 0])
        self.assertEqual(self.client.pr_work_item_ids.call_count, 6)

    def test_strict_threshold_and_both_target_branches(self):
        rows = self.extract(over_sla=True)
        self.assertEqual([row.pr_id for row in rows], [4, 5])
        self.assertEqual([row.is_to_main for row in rows], [False, True])
        self.assertEqual(
            self.client.pr_work_item_ids.call_args_list,
            [call(self.repo, 4), call(self.repo, 5)],
        )

    def test_weekends_and_holidays_affect_threshold(self):
        self.client.pr_last_commit_date.return_value = dt.date(2026, 5, 8)
        self.client.pr_last_approval_date.side_effect = None
        self.client.pr_last_approval_date.return_value = dt.date(2026, 5, 13)
        self.assertEqual(len(self.extract(over_sla=True)), 6)
        report.set_holiday_dates([dt.date(2026, 5, 11)])
        self.assertEqual(self.extract(over_sla=True), [])

    def test_existing_missing_date_fallbacks_are_preserved(self):
        self.client.pr_last_approval_date.side_effect = None
        self.client.pr_last_approval_date.return_value = None
        self.assertEqual([r.business_days for r in self.extract(True)], [9] * 6)
        self.client.pr_last_commit_date.return_value = None
        self.assertEqual(self.extract(True), [])

    def test_holiday_setup_covers_intermediate_years_and_off_resets(self):
        with patch.object(report, "fetch_microsoft_holidays", return_value=[]) as fetch:
            report.configure_holidays(dt.date(2024, 12, 1), dt.date(2026, 1, 1))
        self.assertEqual(fetch.call_args_list, [call(2024), call(2025), call(2026)])
        report.set_holiday_dates([dt.date(2026, 5, 11)])
        report.configure_holidays(self.since, self.until, "off")
        self.assertEqual(
            report.business_days_between(dt.date(2026, 5, 8), dt.date(2026, 5, 13)),
            3,
        )

    def assert_deck(self, path, expected_ids, filtered):
        deck = Presentation(path)
        self.assertEqual(len(deck.slides), 2)
        actual_ids = []
        for slide in deck.slides:
            title = slide.shapes[0].text
            self.assertEqual(">2 business days" in title, filtered)
            table = next(shape.table for shape in slide.shapes if shape.has_table)
            for row in list(table.rows)[1:]:
                actual_ids.append(int(row.cells[1].text.rsplit("/", 1)[1]))
                if filtered:
                    self.assertGreater(int(row.cells[4].text), 2)
        self.assertCountEqual(actual_ids, expected_ids)

    def test_cli_generates_filtered_multimonth_deck_and_normal_report(self):
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(report, "OUTPUT_DIR", Path(tmp)), \
                patch.object(report, "load_config", return_value=([self.repo], self.team)), \
                patch.object(report, "build_default_client", return_value=self.client):
            args = ["--from", "2026-05-01", "--to", "2026-07-31", "--holidays", "off"]
            self.assertEqual(report.main(args + ["--over-sla"]), 0)
            self.assert_deck(
                Path(tmp) / "PR_SLA_2026-05-01_2026-07-31_over_sla.pptx",
                [4, 5], True,
            )
            self.assertEqual(report.main(args), 0)
            self.assert_deck(
                Path(tmp) / "PR_SLA_2026-05-01_2026-07-31.pptx",
                [1, 2, 3, 4, 5, 6], False,
            )
            custom = Path(tmp) / "custom.pptx"
            self.assertEqual(report.main(args + ["--over-sla", "--output", str(custom)]), 0)
            self.assert_deck(custom, [4, 5], True)

    def test_cli_count_only_with_and_without_sla_filter(self):
        all_prs = self.client.completed_prs_for_user.return_value
        for filtered, empty in ((False, False), (True, False), (True, True)):
            with self.subTest(filtered=filtered, empty=empty), tempfile.TemporaryDirectory() as tmp, \
                    contextlib.redirect_stdout(io.StringIO()) as output, \
                    patch.object(report, "OUTPUT_DIR", Path(tmp) / "output"), \
                    patch.object(report, "load_config", return_value=([self.repo], self.team)), \
                    patch.object(report, "build_default_client", return_value=self.client), \
                    patch.object(report, "build_deck") as build:
                self.client.completed_prs_for_user.return_value = [] if empty else all_prs
                args = [
                    "--from", "2026-07-01", "--to", "2026-09-30",
                    "--holidays", "off", "--count-only",
                ]
                if filtered:
                    args.append("--over-sla")
                self.assertEqual(report.main(args), 0)
                expected = 0 if empty else (2 if filtered else 6)
                self.assertIn(f"repo: {expected} PR(s)", output.getvalue())
                self.assertIn(f"Total: {expected} PR(s)", output.getvalue())
                self.assertEqual(">2 business days" in output.getvalue(), filtered)
                build.assert_not_called()
                self.assertFalse((Path(tmp) / "output").exists())
                self.client.completed_prs_for_user.assert_called_with(
                    self.repo, self.team[0][1], dt.date(2026, 7, 1), dt.date(2026, 9, 30),
                )
                report.OUTPUT_DIR.mkdir()
                existing = report.OUTPUT_DIR / "PR_SLA_2026-07-01_2026-09-30.pptx"
                existing.write_bytes(b"existing report")
                self.assertEqual(report.main(args), 0)
                self.assertEqual(existing.read_bytes(), b"existing report")
                self.assertEqual(list(report.OUTPUT_DIR.iterdir()), [existing])
                build.assert_not_called()

    def test_count_only_rejects_output_path(self):
        with contextlib.redirect_stderr(io.StringIO()), \
                patch.object(report, "build_default_client") as client:
            with self.assertRaises(SystemExit) as error:
                report.main(["--count-only", "--output", "unused.pptx"])
            self.assertEqual(error.exception.code, 2)
            client.assert_not_called()

    def test_interactive_decks_and_count_only(self):
        all_prs = self.client.completed_prs_for_user.return_value
        for filtered, empty, count_only in (
            (False, False, False), (True, False, False), (True, True, False),
            (False, False, True), (True, False, True), (True, True, True),
        ):
            with self.subTest(filtered=filtered, empty=empty, count_only=count_only), \
                    tempfile.TemporaryDirectory() as tmp, \
                    contextlib.redirect_stdout(io.StringIO()) as output, \
                    patch.object(interactive, "OUTPUT_DIR", Path(tmp) / "output"), \
                    patch.object(interactive, "CONFIG_XLSX") as config, \
                    patch.object(interactive, "load_config", return_value=([self.repo], self.team)), \
                    patch.object(interactive, "prompt_prerequisites"), \
                    patch.object(interactive, "prompt_confirm_repos"), \
                    patch.object(interactive, "prompt_confirm_team"), \
                    patch.object(interactive, "prompt_auth", return_value=("pat", "test")), \
                    patch.object(interactive, "prompt_month", return_value=(
                        self.since, self.until, "May-July 2026", "2026-05_2026-07",
                    )), \
                    patch.object(interactive, "prompt_sort", return_value=("days", True)) as sort, \
                    patch.object(interactive, "build_deck", wraps=report.build_deck) as build, \
                    patch.object(interactive, "AdoClient", return_value=self.client), \
                    patch.object(interactive, "configure_holidays") as holidays:
                config.exists.return_value = True
                self.client.completed_prs_for_user.return_value = [] if empty else all_prs
                args = ["--over-sla"] if filtered else []
                if count_only:
                    args.append("--count-only")
                self.assertEqual(interactive.main(args), 0)
                if filtered:
                    holidays.assert_called_once_with(self.since, self.until)
                    if empty:
                        self.assertIn("No PRs matched the >2 business-day filter", output.getvalue())
                else:
                    holidays.assert_not_called()
                suffix = "_over_sla" if filtered else ""
                expected_ids = [] if empty else ([4, 5] if filtered else [1, 2, 3, 4, 5, 6])
                self.assertIn(f"Total: {len(expected_ids)} PR(s)", output.getvalue())
                if count_only:
                    build.assert_not_called()
                    sort.assert_not_called()
                    self.assertFalse(interactive.OUTPUT_DIR.exists())
                    self.assertIn("No output file generated", output.getvalue())
                    self.assertNotIn("A deck will still be generated", output.getvalue())
                    continue
                sort.assert_called_once()
                self.assert_deck(
                    interactive.OUTPUT_DIR / f"PR_SLA_2026-05_2026-07{suffix}.pptx",
                    expected_ids, filtered,
                )


if __name__ == "__main__":
    unittest.main()
